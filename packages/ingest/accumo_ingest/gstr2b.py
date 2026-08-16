"""GSTR-2B consolidation / portal extract → inward invoice rows."""

from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from io import BytesIO
from itertools import zip_longest


@dataclass
class TwoBRow:
    supplier_gstin: str
    supplier_name: str
    invoice_number: str
    invoice_date: str
    invoice_value: str
    taxable: str
    igst: str
    cgst: str
    sgst: str
    itc_available: str
    kind: str  # b2b | credit_note


SKIP_SHEETS = {"sheet4", "itc reco", "cgst_tally", "igst_tally", "sgst_tally", "b2ba"}


def _cell(v: object) -> str:
    if v is None:
        return ""
    if hasattr(v, "strftime"):
        return v.strftime("%d/%m/%Y")
    return str(v).strip()


def _dec(raw: str) -> str:
    if not raw:
        return "0"
    try:
        return str(Decimal(str(raw).replace(",", "").replace("₹", "")))
    except InvalidOperation:
        return "0"


def _merge_headers(top: list[str], sub: list[str]) -> list[str]:
    merged: list[str] = []
    for a, b in zip_longest(top, sub, fillvalue=""):
        al, bl = a.lower(), b.lower()
        if bl and (not al or "details" in al):
            merged.append(b)
        else:
            merged.append(a or b)
    return [h.lower() for h in merged]


def _is_header(cells: list[str]) -> bool:
    joined = " ".join(cells).lower()
    return "gstin of supplier" in joined


def _is_subheader(cells: list[str]) -> bool:
    joined = " ".join(cells).lower()
    return "note number" in joined or "invoice number" in joined


def parse_2b_bytes(data: bytes) -> list[TwoBRow]:
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(data), data_only=True)
    rows: list[TwoBRow] = []
    for name in wb.sheetnames:
        if name.lower() in SKIP_SHEETS:
            continue
        kind = "credit_note" if "credit" in name.lower() else "b2b"
        ws = wb[name]
        collected = [[_cell(c) for c in raw] for raw in ws.iter_rows(values_only=True)]
        header_i = next((i for i, cells in enumerate(collected) if _is_header(cells)), None)
        if header_i is None:
            continue
        headers = [c.lower() for c in collected[header_i]]
        data_start = header_i + 1
        if header_i + 1 < len(collected) and _is_subheader(collected[header_i + 1]):
            headers = _merge_headers(collected[header_i], collected[header_i + 1])
            data_start = header_i + 2

        def idx(*names: str) -> int | None:
            for i, h in enumerate(headers):
                if any(n in h for n in names):
                    return i
            return None

        i_gstin = idx("gstin of supplier")
        i_name = idx("trade/legal", "trade", "legal name")
        i_inv = idx("invoice number", "note number")
        i_date = idx("invoice date", "note date")
        i_val = idx("invoice value", "note value")
        i_taxable = idx("taxable value")
        i_igst = idx("integrated tax")
        i_cgst = idx("central tax")
        i_sgst = idx("state/ut tax", "state tax")
        i_itc = idx("itc availability")
        i_type = idx("note type", "invoice type")
        for cells in collected[data_start:]:
            gstin = cells[i_gstin] if i_gstin is not None and i_gstin < len(cells) else ""
            number = cells[i_inv] if i_inv is not None and i_inv < len(cells) else ""
            if not gstin or not number or gstin.lower().startswith("gstin"):
                continue
            if not any(ch.isdigit() for ch in gstin):
                continue
            note_type = cells[i_type] if i_type is not None and i_type < len(cells) else ""
            if kind == "credit_note" and note_type and "debit" in note_type.lower():
                continue
            rows.append(
                TwoBRow(
                    supplier_gstin=gstin,
                    supplier_name=cells[i_name] if i_name is not None and i_name < len(cells) else "",
                    invoice_number=number,
                    invoice_date=cells[i_date] if i_date is not None and i_date < len(cells) else "",
                    invoice_value=_dec(cells[i_val] if i_val is not None and i_val < len(cells) else ""),
                    taxable=_dec(cells[i_taxable] if i_taxable is not None and i_taxable < len(cells) else ""),
                    igst=_dec(cells[i_igst] if i_igst is not None and i_igst < len(cells) else ""),
                    cgst=_dec(cells[i_cgst] if i_cgst is not None and i_cgst < len(cells) else ""),
                    sgst=_dec(cells[i_sgst] if i_sgst is not None and i_sgst < len(cells) else ""),
                    itc_available=cells[i_itc] if i_itc is not None and i_itc < len(cells) else "",
                    kind=kind,
                )
            )
    wb.close()
    return rows
