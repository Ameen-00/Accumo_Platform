"""Turn a CSV or xlsx upload into headers + row dicts. No database."""

from __future__ import annotations

import csv
import io
from pathlib import Path


def _cells(row: list[object]) -> list[str]:
    out = []
    for cell in row:
        if cell is None:
            out.append("")
        else:
            out.append(str(cell).strip())
    return out


def _first_header_index(rows: list[list[str]]) -> int:
    for i, row in enumerate(rows[:10]):
        nonempty = [c for c in row if c]
        if len(nonempty) >= 2:
            return i
    return 0


def parse_csv_bytes(data: bytes) -> tuple[list[str], list[dict[str, str]]]:
    text = data.decode("utf-8-sig")
    reader = csv.reader(io.StringIO(text))
    raw = [_cells(list(r)) for r in reader]
    return _tabulate(raw)


def parse_xlsx_bytes(data: bytes) -> tuple[list[str], list[dict[str, str]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("xlsx support needs openpyxl — save as CSV or pip install openpyxl") from exc
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    raw = []
    for row in ws.iter_rows(values_only=True):
        raw.append(_cells(list(row or [])))
    wb.close()
    return _tabulate(raw)


def parse_upload(filename: str, data: bytes) -> tuple[list[str], list[dict[str, str]]]:
    name = filename.lower()
    if name.endswith(".xlsx") or name.endswith(".xls"):
        return parse_xlsx_bytes(data)
    return parse_csv_bytes(data)


def parse_path(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    return parse_upload(path.name, path.read_bytes())


def _tabulate(raw: list[list[str]]) -> tuple[list[str], list[dict[str, str]]]:
    raw = [r for r in raw if any(c for c in r)]
    if not raw:
        return [], []
    head_i = _first_header_index(raw)
    headers = raw[head_i]
    # Drop empty trailing header cells
    while headers and headers[-1] == "":
        headers = headers[:-1]
    rows: list[dict[str, str]] = []
    for line in raw[head_i + 1 :]:
        line = (line + [""] * len(headers))[: len(headers)]
        if not any(line):
            continue
        rows.append({headers[i]: line[i] for i in range(len(headers))})
    return headers, rows
