"""HDFC-style account statement → payment rows.

Indian retail statements put the header after a block of address lines.
We never invent a vendor from a narration — the raw narration is kept.
"""

from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from io import BytesIO


@dataclass
class BankLine:
    payment_date: str
    narration: str
    reference: str
    amount: str
    direction: str  # out | in
    balance: str


def _cell(v: object) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    return str(v).strip()


def _rows_from_xls(data: bytes) -> list[list[str]]:
    try:
        import xlrd
    except ImportError as exc:
        raise RuntimeError("legacy .xls needs xlrd") from exc
    book = xlrd.open_workbook(file_contents=data)
    sheet = book.sheet_by_index(0)
    return [[_cell(sheet.cell_value(r, c)) for c in range(sheet.ncols)] for r in range(sheet.nrows)]


def _rows_from_xlsx(data: bytes) -> list[list[str]]:
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(data), data_only=True, read_only=True)
    ws = wb.active
    out = []
    for row in ws.iter_rows(values_only=True):
        out.append([_cell(c) for c in (row or [])])
    wb.close()
    return out


def _header_index(rows: list[list[str]]) -> int | None:
    for i, row in enumerate(rows):
        joined = " ".join(c.lower() for c in row)
        if "narration" in joined and "date" in joined:
            return i
    return None


def parse_bank_bytes(data: bytes, filename: str) -> list[BankLine]:
    name = filename.lower()
    rows = _rows_from_xls(data) if name.endswith(".xls") and not name.endswith(".xlsx") else _rows_from_xlsx(data)
    head_i = _header_index(rows)
    if head_i is None:
        return []
    header = [c.lower() for c in rows[head_i]]

    def col(*names: str) -> int | None:
        for i, h in enumerate(header):
            if any(n in h for n in names):
                return i
        return None

    i_date = col("date")
    i_narr = col("narration", "description", "particular")
    i_ref = col("ref", "chq", "cheque")
    i_wd = col("withdrawal", "debit")
    i_dep = col("deposit", "credit")
    i_bal = col("balance")
    out: list[BankLine] = []
    for row in rows[head_i + 1 :]:
        if not row or (row[0].startswith("*") if row else False):
            continue
        date = row[i_date] if i_date is not None and i_date < len(row) else ""
        if not date or not any(ch.isdigit() for ch in date):
            continue
        if "/" not in date and "-" not in date:
            continue
        narr = row[i_narr] if i_narr is not None and i_narr < len(row) else ""
        wd = row[i_wd] if i_wd is not None and i_wd < len(row) else ""
        dep = row[i_dep] if i_dep is not None and i_dep < len(row) else ""
        try:
            if wd:
                amt, direction = str(Decimal(str(wd).replace(",", ""))), "out"
            elif dep:
                amt, direction = str(Decimal(str(dep).replace(",", ""))), "in"
            else:
                continue
        except (InvalidOperation, ValueError):
            continue
        out.append(
            BankLine(
                payment_date=date,
                narration=narr,
                reference=row[i_ref] if i_ref is not None and i_ref < len(row) else "",
                amount=amt,
                direction=direction,
                balance=row[i_bal] if i_bal is not None and i_bal < len(row) else "",
            )
        )
    return out
